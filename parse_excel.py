#!/usr/bin/env python3
"""Parse AP_EAPCET_Colleges_Details.xlsx and output structure."""
import openpyxl
import json
import os

base = os.path.dirname(os.path.abspath(__file__))
path = os.path.join(base, "AP_EAPCET_Colleges_Details.xlsx")
out = os.path.join(base, "xlsx_parsed.txt")

if not os.path.exists(path):
    print(f"File not found: {path}")
    exit(1)

wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
lines = []
lines.append("SHEETS: " + str(wb.sheetnames))
lines.append("")

for sn in wb.sheetnames[:2]:
    ws = wb[sn]
    lines.append(f"--- Sheet: {sn} ---")
    for i, row in enumerate(ws.iter_rows(max_row=50, values_only=True)):
        lines.append(f"{i}: {row}")
    lines.append("")

wb.close()

with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

print(f"Wrote {len(lines)} lines to {out}")
